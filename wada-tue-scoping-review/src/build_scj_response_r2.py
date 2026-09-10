#!/usr/bin/env python3
"""Build point-by-point Response to Reviewers for SCJ-D-26-00070R1
minor revision (second round)."""

import os
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from create_scj_review_part1 import setup_styles, add_heading_styled

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(BASE_DIR, "manuscripts", "SCJ_Response_to_Reviewers_R2.docx")

doc = Document()
setup_styles(doc)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Response to Reviewers")
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    "SCJ-D-26-00070R1: Mind the Gap: A Scoping Review of Discrepancies "
    "Between WADA Therapeutic Use Exemption and Current Clinical Practice "
    "Guidelines"
)
run.font.size = Pt(12)
run.italic = True

doc.add_paragraph()


def add_para(text, size=12, bold=False, italic=False, indent=None):
    pr = doc.add_paragraph()
    if indent:
        pr.paragraph_format.left_indent = Cm(indent)
    run = pr.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return pr


def add_comment(text):
    pr = doc.add_paragraph()
    pr.paragraph_format.left_indent = Cm(0.5)
    run = pr.add_run("Comment: ")
    run.bold = True
    run.font.size = Pt(11)
    run = pr.add_run(text)
    run.italic = True
    run.font.size = Pt(11)


def add_response(text):
    pr = doc.add_paragraph()
    pr.paragraph_format.left_indent = Cm(0.5)
    run = pr.add_run("Response: ")
    run.bold = True
    run.font.size = Pt(11)
    run = pr.add_run(text)
    run.font.size = Pt(11)


add_heading_styled(doc, "General Response", 1)
add_para(
    "We thank the Editor and both reviewers for their continued careful "
    "reading of our manuscript. In this minor revision we have made only the "
    "changes needed to address the comments: the keywords have been replaced "
    "with terms that do not appear in the title, no sentence now begins with "
    "an acronym, all abbreviations are established at first use in both the "
    "abstract and the paper, a brief methods description has been added "
    "explaining how the gap severity ratings in Figure 2 were derived, and "
    "each of Reviewer #2's line-level corrections has been applied. A clean "
    "revised manuscript and a marked manuscript (deletions shown in blue "
    "strikethrough, insertions shown in red) are provided."
)

doc.add_paragraph()
add_heading_styled(doc, "Editor Comments", 1)

items_editor = [
    (
        "key word change please use words not in title",
        "The keywords have been replaced with terms that do not appear in the "
        "title: stimulant medications; exercise-induced bronchoconstriction; "
        "glucocorticoids; hypogonadism; sports medicine; medication "
        "management.",
    ),
    (
        "DO NOT begin a sentence with an acronym.",
        "All sentences that began with an acronym now begin with the full "
        "expansion (e.g., \u201cThe World Anti-Doping Agency\u2026\u201d, "
        "\u201cThe Global Initiative for Asthma\u2026\u201d, "
        "\u201cGlucagon-like peptide-1 receptor agonists\u2026\u201d, "
        "\u201cStrength and conditioning professionals\u2026\u201d).",
    ),
    (
        "Please make sure that all abbreviations are established with first "
        "use in Both abstract (e.g. GLP) and paper itself.",
        "In the abstract, glucagon-like peptide-1 (GLP-1) is now defined at "
        "first use and the PRISMA-ScR reporting guideline is now spelled out. "
        "In the body, the following abbreviations are now established at "
        "first use: WADA, S&C, EAU, attention-deficit/hyperactivity disorder "
        "(ADHD), GLP-1, GIP, FDA, EMA, PMDA, SGLT2, ACE, ARNi, WHO, PMOS "
        "(section heading), and DSM-5.",
    ),
]
for c, r in items_editor:
    add_comment(c)
    add_response(r)
    doc.add_paragraph()

add_heading_styled(doc, "Reviewer #1", 1)
add_comment(
    "The authors have addressed most of my comments. I only have one more - "
    "please add methods about the gap sensitivity in Figure 2. How did you do "
    "this? It is not mentioned."
)
add_response(
    "A methods description has been added at the end of the Data Charting "
    "Process section (Methods): \u201cFor each disease area, scores were "
    "assigned for each dimension through a qualitative synthesis of the "
    "charted evidence, with disagreements resolved by discussion and "
    "consensus; the resulting matrix is presented as a heat map (Figure "
    "2).\u201d"
)

doc.add_paragraph()
add_heading_styled(doc, "Reviewer #2", 1)
add_para(
    "We thank Reviewer #2 for confirming the improvements and for the "
    "detailed line-level corrections. Each item is addressed below.",
    size=11,
)
doc.add_paragraph()

items_r2 = [
    (
        "Page 4, Line 11: Replace \u201cIfs\u201d with \u201cIFs\u201d",
        "Corrected: \u201cinternational federations (IFs)\u201d.",
    ),
    (
        "Page 9, Line 33: Recommend rephrasing \u201colder recreational "
        "athletes\u201d as this could be offensive to competitive masters "
        "athletes who are in fact competitive and elite by age standards.",
        "Rephrased to \u201cyounger elite competitors and masters athletes "
        "alike\u201d.",
    ),
    (
        "Page 11, Line 33: Delete extra space between \u201c2022\u201d and "
        "\u201cdesignate\u201d",
        "The extra space has been deleted.",
    ),
    (
        "Page 11, Line 42: Appears to be one or more references missing from "
        "empty parentheses at the end of the sentence",
        "The sentence now carries its supporting references: "
        "\u201c\u2026highest response rates among available treatments (22, "
        "4).\u201d We also audited every parenthetical citation in the "
        "manuscript to confirm that no parentheses are left empty.",
    ),
    (
        "Page 12, Lines 4,7: The following could be considered a "
        "mischaracterization: \u201c\u2026yet WADA's regulatory framework "
        "treats it as if pharmacotherapy can be safely interrupted for "
        "competitive events.\u201d WADA makes no such suggestion - the rule "
        "is in place because the use of stimulants has the potential to be "
        "performance enhancing.",
        "Agreed; we have removed the mischaracterization. The sentence now "
        "reads: \u201cThis creates a fundamental paradox: ADHD is a "
        "continuous neurodevelopmental condition that does not remit on "
        "competition days, yet first-line stimulant pharmacotherapy is "
        "prohibited in-competition.\u201d",
    ),
    (
        "Page 13, Line 9: Replace \u201cIFs\u201d with \u201cIF\u201d",
        "Corrected: \u201cpractical implementation varies by NADO and "
        "IF\u201d.",
    ),
    (
        "Page 13, Line 50: After \u201cEuropean Medicines Agency\u201d insert "
        "\u201c(EMA)\u201d since that acronym follows in the next sentence",
        "Inserted \u201c(EMA)\u201d after \u201cEuropean Medicines "
        "Agency\u201d. \u201c(FDA)\u201d and \u201c(PMDA)\u201d were added in "
        "the same way so that all three agency abbreviations are established "
        "at first use.",
    ),
    (
        "Page 24, Line 20: Delete \u201cSociety\u201d as it is a repeated "
        "word following \u201cSociety's\u201d",
        "Corrected: \u201c\u2026in light of the Endocrine Society\u2019s "
        "symptom-plus-biochemistry approach\u2026\u201d.",
    ),
]
for c, r in items_r2:
    add_comment(c)
    add_response(r)
    doc.add_paragraph()

add_para(
    "We hope the manuscript is now suitable for publication and thank the "
    "editor and reviewers for their time and constructive feedback.",
    size=11,
)

doc.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
