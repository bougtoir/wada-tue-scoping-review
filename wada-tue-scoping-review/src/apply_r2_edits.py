#!/usr/bin/env python3
"""Apply R2 (minor revision) edits to the SCJ main body docx.

Edits are minimal text replacements mapped to the SCJ-D-26-00070R1
decision letter comments. Paragraphs are rebuilt run-by-run preserving
the Times New Roman 12pt style and the font-superscript convention for
in-text citation numbers.
"""
import copy
import os
import re
import shutil
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAINBODY = os.path.join(BASE_DIR, "manuscripts", "SCJ_Narrative_Review_mainbody.docx")
TITLE_PAGE = os.path.join(BASE_DIR, "manuscripts", "SCJ_Title_Page.docx")

WADA_EXP = "The World Anti-Doping Agency"

NEW_KEYWORDS = (
    "Keywords: stimulant medications; exercise-induced bronchoconstriction; "
    "glucocorticoids; hypogonadism; sports medicine; medication management"
)

# paragraph index -> list of (old, new) exact-substring replacements
EDITS = {
    2: [(
        "Keywords: therapeutic use exemption; anti-doping; clinical practice "
        "guidelines; scoping review; prohibited list; athlete health",
        NEW_KEYWORDS,
    )],
    4: [
        (
            "This PRISMA-ScR\u2013compliant scoping review systematically mapped",
            "This scoping review, conducted in accordance with the Preferred "
            "Reporting Items for Systematic Reviews and Meta-Analyses "
            "extension for Scoping Reviews (PRISMA-ScR), systematically mapped",
        ),
        (
            "emerging GLP-1 receptor agonist therapies",
            "emerging glucagon-like peptide-1 (GLP-1) receptor agonist therapies",
        ),
    ],
    6: [
        (
            "on the WADA Prohibited List",
            "on the World Anti-Doping Agency (WADA) Prohibited List",
        ),
        (
            "WADA's stated objective is to permit",
            f"{WADA_EXP}'s stated objective is to permit",
        ),
    ],
    7: [(
        "WADA maintains the Prohibited List",
        f"{WADA_EXP} maintains the Prohibited List",
    )],
    8: [(
        "including S&C practitioners",
        "including strength and conditioning (S&C) practitioners",
    )],
    11: [(
        "The role of S&C professionals in this context",
        "The role of strength and conditioning professionals in this context",
    )],
    12: [("international federations (Ifs)", "international federations (IFs)")],
    25: [(
        "across five dimensions: first-line treatment prohibition, TUE process "
        "complexity, guideline update lag, athlete impact, and alternative "
        "treatment availability.",
        "across five dimensions: first-line treatment prohibition, TUE process "
        "complexity, guideline update lag, athlete impact, and alternative "
        "treatment availability. For each disease area, scores were assigned "
        "for each dimension through a qualitative synthesis of the charted "
        "evidence, with disagreements resolved by discussion and consensus; "
        "the resulting matrix is presented as a heat map (Figure 2).",
    )],
    27: [(
        "Endocrine Society and EAU guidelines for hypogonadism",
        "Endocrine Society and European Association of Urology (EAU) guidelines "
        "for hypogonadism",
    )],
    32: [
        (
            "the prevalence of ADHD among elite athletes",
            "the prevalence of attention-deficit/hyperactivity disorder (ADHD) "
            "among elite athletes",
        ),
        (
            "younger elite competitors and older recreational athletes alike",
            "younger elite competitors and masters athletes alike",
        ),
    ],
    34: [(
        "emerging gaps related to GLP-1 receptor agonists",
        "emerging gaps related to glucagon-like peptide-1 (GLP-1) receptor "
        "agonists",
    )],
    42: [
        ("2022  designate", "2022 designate"),
        (
            "highest response rates among available treatments.",
            "highest response rates among available treatments (22, 4).",
        ),
    ],
    43: [
        (
            "WADA's rationale for the in-competition ban is that",
            f"{WADA_EXP}'s rationale for the in-competition ban is that",
        ),
        (
            "yet WADA's regulatory framework treats it as if pharmacotherapy "
            "can be safely interrupted for competitive events.",
            "yet first-line stimulant pharmacotherapy is prohibited "
            "in-competition.",
        ),
    ],
    44: [(
        "per DSM-5 criteria",
        "per the Diagnostic and Statistical Manual of Mental Disorders, "
        "Fifth Edition (DSM-5) criteria",
    )],
    46: [
        (
            "TUE access for stimulant medications is not uniformly available",
            "Access to TUEs for stimulant medications is not uniformly available",
        ),
        ("varies by NADO and IFs.", "varies by NADO and IF."),
    ],
    48: [(
        "dual GIP/GLP-1 agonists (tirzepatide)",
        "dual glucose-dependent insulinotropic polypeptide (GIP)/GLP-1 "
        "agonists (tirzepatide)",
    )],
    49: [
        (
            "US Food and Drug Administration approval in December 2017",
            "US Food and Drug Administration (FDA) approval in December 2017",
        ),
        (
            "European Medicines Agency authorization in February 2018",
            "European Medicines Agency (EMA) authorization in February 2018",
        ),
        (
            "Pharmaceuticals and Medical Devices Agency approval in March 2018",
            "Pharmaceuticals and Medical Devices Agency (PMDA) approval in "
            "March 2018",
        ),
    ],
    50: [
        (
            "WADA's rationale relates to their weight-loss properties",
            f"{WADA_EXP}'s rationale relates to their weight-loss properties",
        ),
        (
            "metformin and SGLT2 inhibitors",
            "metformin and sodium-glucose cotransporter-2 (SGLT2) inhibitors",
        ),
    ],
    51: [(
        "GLP-1 receptor agonists reduce appetite",
        "Glucagon-like peptide-1 receptor agonists reduce appetite",
    )],
    55: [(
        "WADA's concern is that TRT can restore",
        f"{WADA_EXP}'s concern is that TRT can restore",
    )],
    62: [(
        "ACE inhibitors/ARNi",
        "angiotensin-converting enzyme (ACE) inhibitors/angiotensin "
        "receptor\u2013neprilysin inhibitors (ARNi)",
    )],
    63: [(
        "WADA's rationale for prohibiting beta-blockers in these sports",
        f"{WADA_EXP}'s rationale for prohibiting beta-blockers in these sports",
    )],
    64: [(
        "PMOS and Female Fertility Treatment",
        "Polyendocrine Metabolic Ovarian Syndrome (PMOS) and Female "
        "Fertility Treatment",
    )],
    65: [(
        "to polyendocrine metabolic ovarian syndrome (PMOS) to better reflect",
        "to polyendocrine metabolic ovarian syndrome to better reflect",
    )],
    66: [
        ("and WHO recommendations", "and World Health Organization (WHO) recommendations"),
        (
            "WADA's rationale is that aromatase inhibitors",
            f"{WADA_EXP}'s rationale is that aromatase inhibitors",
        ),
    ],
    70: [
        ("WADA has developed tools", f"{WADA_EXP} has developed tools"),
        (
            "GINA has published annual updates",
            "The Global Initiative for Asthma has published annual updates",
        ),
    ],
    75: [(
        "WADA's mandate is to protect clean sport",
        f"{WADA_EXP}'s mandate is to protect clean sport",
    )],
    78: [(
        "S&C professionals, while not prescribers",
        "Strength and conditioning professionals, while not prescribers",
    )],
    83: [(
        "Endocrine Society\u2019s Society symptom-plus-biochemistry",
        "Endocrine Society\u2019s symptom-plus-biochemistry",
    )],
    91: [(
        "S&C professionals are often the first",
        "Strength and conditioning professionals are often the first",
    )],
}

# Title-page edits: keywords line + title singular ("Exemption", per EM record)
TITLE_PAGE_EDITS = {
    0: [(
        "WADA Therapeutic Use Exemptions and Current",
        "WADA Therapeutic Use Exemption and Current",
    )],
    12: [(
        "WADA Therapeutic Use Exemptions and Current",
        "WADA Therapeutic Use Exemption and Current",
    )],
    16: [(
        "Keywords: therapeutic use exemption; anti-doping; clinical practice "
        "guidelines; scoping review; prohibited list; athlete health",
        NEW_KEYWORDS,
    )],
}

CITATION_RE = re.compile(r"\((\d+(?:\s*,\s*\d+)*)\)")
NUM_RUN_RE = re.compile(r"(\d+)")


def _style_run(r, size=12, bold=False, color=None, strike=False):
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    if strike:
        r.font.strike = True
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    return r


def _add_citation_aware_runs(p, text, size=12, bold=False):
    """Add runs so citation numbers inside (N, N) render as superscript."""
    pos = 0
    for m in CITATION_RE.finditer(text):
        if m.start() > pos:
            _style_run(p.add_run(text[pos:m.start()]), size, bold)
        _style_run(p.add_run("("), size, bold)
        inner = m.group(1)
        for j, part in enumerate(NUM_RUN_RE.split(inner)):
            if not part:
                continue
            r = _style_run(p.add_run(part), size, bold)
            if part.isdigit():
                r.font.superscript = True
        _style_run(p.add_run(")"), size, bold)
        pos = m.end()
    if pos < len(text):
        _style_run(p.add_run(text[pos:]), size, bold)


def _clear_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def _para_is_bold(p):
    return any(r.font.bold for r in p.runs)


def _rebuild_paragraph(p, new_text):
    bold = _para_is_bold(p)
    _clear_runs(p)
    _add_citation_aware_runs(p, new_text, bold=bold)


def apply_edits(path, edits):
    doc = Document(path)
    old_texts = {}
    for idx, repls in sorted(edits.items()):
        p = doc.paragraphs[idx]
        text = p.text
        old_texts[idx] = text
        for old, new in repls:
            if old not in text:
                raise SystemExit(
                    f"para {idx}: pattern not found:\n  {old!r}\nin:\n  {text!r}"
                )
            text = text.replace(old, new, 1)
        _rebuild_paragraph(p, text)
    doc.save(path)
    return old_texts


def main():
    shutil.copy2(MAINBODY, "/tmp/mainbody_r1_base.docx")
    old = apply_edits(MAINBODY, EDITS)
    print(f"mainbody: edited {len(old)} paragraphs")
    apply_edits(TITLE_PAGE, TITLE_PAGE_EDITS)
    print("title page: edited")


if __name__ == "__main__":
    main()
